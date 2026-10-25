from pathlib import Path
import json, shutil, textwrap
import pandas as pd
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

root=Path(__file__).resolve().parent
reports=root/'reports'; figures=reports/'figures'; figures.mkdir(parents=True,exist_ok=True)
metrics=json.loads((reports/'ml_metrics.json').read_text(encoding='utf-8'))
rf=metrics['random_forest_mae']; base=metrics['previous_week_baseline_mae']; improvement=metrics['mae_improvement_percent_vs_baseline']

# MAE comparison chart
plt.figure(figsize=(8,4.8))
plt.bar(['Random Forest','Previous-week baseline'],[rf,base])
plt.ylabel('MAE (orders; lower is better)')
plt.title('Temporal holdout: weeks 136–145')
plt.tight_layout()
plt.savefig(figures/'ml_mae_comparison.png',dpi=180,bbox_inches='tight')
plt.close()

# Update README
readme=(root/'README.md').read_text(encoding='utf-8')
if '## التنبؤ بالطلب (إضافة تعلّم الآلة)' not in readme:
 readme += f'''\n\n## التنبؤ بالطلب (إضافة تعلّم الآلة)\n\nشغّل من مجلد المشروع الرئيسي:\n\n```bash\npython run_ml_analysis.py\n```\n\nينفذ البرنامج تقييمًا زمنيًا على آخر 10 أسابيع معلومة (136–145)، ويقارن نموذج Random Forest بخط أساس يعتمد على طلب الأسبوع السابق. لتناسب سرعة التنفيذ الحواسيب العادية، تُستخدم عينة عشوائية ثابتة تصل إلى 120,000 صف أثناء التقييم و150,000 صف عند تدريب النموذج النهائي.\n\n**نتيجة التقييم الحالي:** MAE نموذج Random Forest = **{rf:.3f}** طلب، وMAE خط الأساس = **{base:.3f}** طلب. التغير في MAE مقارنة بخط الأساس = **{improvement:.2f}%**؛ والقيمة السالبة تعني أن النموذج لم يتفوق على خط الأساس في هذا الاختبار. لذلك لا ينبغي وصف نموذج الغابة العشوائية بأنه أفضل حاليًا؛ خط الأساس كان أدق على فترة التحقق هذه.\n\nالملفات الناتجة:\n- `reports/ml_metrics.json` و`reports/ml_metrics.csv`: نتائج المقارنة.\n- `reports/ml_validation_predictions.csv`: القيم الفعلية وتنبؤات النموذج وخط الأساس لأسابيع التحقق.\n- `reports/test_predictions.csv`: تنبؤات `num_orders` لصفوف `test.csv`، مع الأعمدة `id` و`num_orders`.\n- `reports/ml_feature_importance.csv`: أهمية الميزات وفق النموذج المدرب.\n- `reports/figures/ml_mae_comparison.png`: رسم مقارنة MAE.\n\n**تنبيه منهجي:** لا يحتوي `test.csv` على `num_orders` الفعلي، لذا لا يمكن حساب MAE عليه. التقييم العادل تم على أسابيع من بيانات التدريب محجوبة عن النموذج. كما أن بيانات الطلب تاريخية رصدية، والارتباطات لا تثبت السببية.\n'''
 (root/'README.md').write_text(readme,encoding='utf-8')

# Append to client report
report=reports/'client_report.md'
text=report.read_text(encoding='utf-8')
if '## إضافة اختيارية: التنبؤ بالطلب وتقييم النموذج' not in text:
 text += f'''\n\n## إضافة اختيارية: التنبؤ بالطلب وتقييم النموذج\n\n### المنهجية\nتم حجز الأسابيع العشرة الأخيرة من بيانات التدريب (136–145) للاختبار الزمني. دُرّب نموذج Random Forest على الأسابيع السابقة فقط، مع تحويل الهدف باستخدام `log1p` ثم إعادة التنبؤ إلى مقياس عدد الطلبات. ولتوفير وقت التنفيذ على الحواسيب الشخصية، استخدم التدريب عينة عشوائية ثابتة تصل إلى 120,000 صف في التقييم، ثم عينة تصل إلى 150,000 صف لتدريب النموذج المستخدم في `test.csv`.\n\n### النتائج\n- MAE لنموذج Random Forest: **{rf:.3f}** طلب.\n- MAE لخط الأساس (طلب الأسبوع السابق لنفس المركز والوجبة، مع استخدام متوسطات بديلة عند غياب السجل): **{base:.3f}** طلب.\n- تغير MAE مقارنة بخط الأساس: **{improvement:.2f}%**.\n\n**التفسير:** في هذا الاختبار الزمني، كان خط الأساس أفضل من نموذج Random Forest؛ لأن MAE الأقل أفضل. هذه نتيجة مهمة وليست فشلًا: فهي تمنعنا من الادعاء بأن نموذج التعلم الآلي يحسن التنبؤ قبل أن يثبت ذلك. يمكن لاحقًا تجربة ميزات زمنية إضافية، أو ضبط النموذج، أو اختبار نموذج آخر على نفس تقسيم الأسابيع.\n\n### تنبؤات ملف الاختبار\nأُنشئ الملف `reports/test_predictions.csv` ويحتوي على `id` و`num_orders` المتنبأ بهما لعدد **{sum(1 for _ in open(root/'data/raw/test.csv',encoding='utf-8'))-1:,}** صفًا من `test.csv`. هذه تنبؤات وليست أعدادًا فعلية مؤكدة. ولأن `test.csv` لا يحتوي على الهدف الحقيقي، لا يمكن حساب MAE لهذا الملف.\n\nالرسم `reports/figures/ml_mae_comparison.png` يعرض مقارنة الخطأ بين الطريقتين.\n'''
 report.write_text(text,encoding='utf-8')

# Use the richer Arabic defense deck as the project deck, then append two ML slides.
source=Path('/mnt/data/عرض_مناقشة_مشروع_تحليل_الطلب_على_الوجبات.pptx')
target=reports/'final_presentation.pptx'
if source.exists() and source.resolve()!=target.resolve():
 shutil.copy2(source,target)
prs=Presentation(target)
# Append slide helper using the presentation's existing layout style
layout=prs.slide_layouts[5] if len(prs.slide_layouts)>5 else prs.slide_layouts[1]

def add_slide(title, body, chart_path=None):
 slide=prs.slides.add_slide(layout)
 # title placeholder if available
 title_shape=None
 for sh in slide.shapes:
  if sh.is_placeholder and sh.placeholder_format.type==1:
   title_shape=sh; break
 if title_shape is None:
  title_shape=slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.7))
 title_shape.text=title
 for p in title_shape.text_frame.paragraphs:
  p.font.name='Arial'; p.font.size=Pt(27); p.font.bold=True; p.font.color.rgb=RGBColor(31,78,121)
 body_box=slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(6.1 if chart_path else 11.5), Inches(5.2))
 tf=body_box.text_frame; tf.word_wrap=True
 for i,line in enumerate(body.split('\n')):
  para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
  para.text=line; para.font.name='Arial'; para.font.size=Pt(19 if len(body.split('\n'))<6 else 16)
  para.space_after=Pt(12); para.alignment=PP_ALIGN.RIGHT
 if chart_path and Path(chart_path).exists():
  slide.shapes.add_picture(str(chart_path), Inches(7.0), Inches(1.7), width=Inches(5.6))

add_slide('إضافة تعلّم الآلة: كيف قيّمنا التنبؤ؟',
 f'• استخدمنا Random Forest للتنبؤ بعدد الطلبات.\n• درّبنا النموذج على الأسابيع السابقة، وحجزنا الأسابيع 136–145 للتقييم.\n• قارنا النموذج بخط أساس يستخدم طلب الأسبوع السابق للمركز والوجبة.\n• استخدمنا MAE: كلما كانت القيمة أقل كان الخطأ أصغر.\n• استخدمنا عينة تدريب ثابتة لتقليل وقت التنفيذ على الحاسوب الشخصي.')
add_slide('نتيجة التنبؤ: خط الأساس أدق في الاختبار الحالي',
 f'• MAE نموذج Random Forest: {rf:.2f} طلب.\n• MAE خط الأساس: {base:.2f} طلب.\n• التغير مقابل خط الأساس: {improvement:.2f}% (السالب يعني أن النموذج أسوأ).\n• النتيجة: لا نعتمد نموذج الغابة العشوائية بوصفه أفضل قبل تحسينه وإعادة تقييمه.\n• ملف test_predictions.csv يحتوي تنبؤات فقط؛ لا يمكن تقييمه دون أعداد فعلية.',
 figures/'ml_mae_comparison.png')
prs.save(target)

# Also save a separate copy of the updated Arabic defense deck at /mnt/data for easy access.
shutil.copy2(target,Path('/mnt/data/عرض_مناقشة_مشروع_تحليل_الطلب_على_الوجبات_محدث.pptx'))
print('Updated README, report, chart, and Arabic presentation.')
print(f'MAE Random Forest={rf:.3f}; baseline={base:.3f}; delta={improvement:.2f}%')
